"""
CLONACION_CARPETA — publicar_inventario_sheets.py
-------------------------------------------------
Sube el Excel de inventario a UNA Google Sheet fija (se sobrescribe cada corrida).
Guarda el ID en hoja_inventario_id.txt para reutilizar el mismo enlace.
Comparte la hoja con los correos de CORREOS_AVISO.
Lo llama clone_carpeta_drive.py; el link resultante va en el correo de clonación.
"""

# ---------------------------------------------------------------------------
# IMPORTS
# ---------------------------------------------------------------------------

from __future__ import annotations

# Logs de permisos / avisos.
import logging

# Ruta del Excel y del archivo hoja_inventario_id.txt.
from pathlib import Path

# Credenciales OAuth del clon.
from google.oauth2.credentials import Credentials

# APIs Sheets (escribir) y Drive (compartir).
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

# Leer el Excel local hoja por hoja.
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# Archivo local donde se guarda el ID de la Sheet (NO subir a Git).
ID_FILE = "hoja_inventario_id.txt"
# Título de la hoja en Google Drive.
TITULO_HOJA = "Inventario clonación Drive"


def _leer_id(base: Path) -> str:
    """# Aquí se lee el ID guardado de la Sheet (si ya existía de otra corrida)."""
    ruta = base / ID_FILE
    if not ruta.is_file():
        return ""
    return ruta.read_text(encoding="utf-8").strip()


def _guardar_id(base: Path, sheet_id: str) -> None:
    """# Aquí se guarda el ID para reutilizar la MISMA Sheet la próxima vez."""
    (base / ID_FILE).write_text(sheet_id, encoding="utf-8")


def _filas_hoja(ws) -> list[list[str]]:
    """# Aquí se convierte una hoja Excel a filas de texto para Sheets."""
    out: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if c is None else str(c) for c in row])
    return out or [[""]]


def _asegurar_pestanas(sheets, spreadsheet_id: str, titulos: list[str]) -> None:
    """# Aquí se crean pestañas faltantes si el Excel trae hojas nuevas."""
    meta = sheets.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    existentes = {s["properties"]["title"] for s in meta.get("sheets", [])}
    pedidos = []
    for t in titulos:
        if t not in existentes:
            pedidos.append({"addSheet": {"properties": {"title": t}}})
    if pedidos:
        sheets.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id, body={"requests": pedidos}
        ).execute()


def _rgb(hex_color: str) -> dict:
    """# Aquí se convierte un color #RRGGBB al formato RGB 0–1 de Sheets."""
    h = hex_color.lstrip("#")
    return {
        "red": int(h[0:2], 16) / 255,
        "green": int(h[2:4], 16) / 255,
        "blue": int(h[4:6], 16) / 255,
    }


def _aplicar_formato(sheets, spreadsheet_id: str, titulos: list[str], n_cols: dict[str, int]) -> None:
    """# Aquí se aplica formato (título azul, cabeceras, anchos) parecido al Excel."""
    meta = sheets.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta.get("sheets", [])}
    titulo_c = _rgb("1F4E79")
    cab_c = _rgb("D6EAF8")
    pedidos = []
    anchos = {
        "Cómo leer": [200, 720],
        "Resumen": [140, 220, 220, 120, 120, 120, 120, 160, 160, 200],
        "Temas y materiales": [120, 220, 220, 280] + [90] * 20,
        "Faltantes": [90, 120, 220, 220, 160, 280, 120, 120, 320],
        "Archivos faltantes": [140, 520, 280],
        "Detalle por carpeta": [120, 420, 220, 110, 110, 130, 130, 140, 140, 160, 220],
    }
    for titulo in titulos:
        sid = ids.get(titulo)
        if sid is None:
            continue
        cols = max(n_cols.get(titulo, 2), 2)
        fila_cab = 1 if titulo == "Cómo leer" else 2
        pedidos.append(
            {
                "repeatCell": {
                    "range": {
                        "sheetId": sid,
                        "startRowIndex": 0,
                        "endRowIndex": 1,
                        "startColumnIndex": 0,
                        "endColumnIndex": cols,
                    },
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": titulo_c,
                            "textFormat": {
                                "foregroundColor": {"red": 1, "green": 1, "blue": 1},
                                "bold": True,
                                "fontFamily": "Calibri",
                                "fontSize": 11,
                            },
                        }
                    },
                    "fields": "userEnteredFormat(backgroundColor,textFormat)",
                }
            }
        )
        pedidos.append(
            {
                "repeatCell": {
                    "range": {
                        "sheetId": sid,
                        "startRowIndex": fila_cab,
                        "endRowIndex": fila_cab + 1,
                        "startColumnIndex": 0,
                        "endColumnIndex": cols,
                    },
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": cab_c,
                            "textFormat": {"bold": True, "fontFamily": "Calibri", "fontSize": 11},
                            "horizontalAlignment": "CENTER",
                            "wrapStrategy": "WRAP",
                        }
                    },
                    "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,wrapStrategy)",
                }
            }
        )
        pedidos.append(
            {
                "updateSheetProperties": {
                    "properties": {
                        "sheetId": sid,
                        "gridProperties": {"frozenRowCount": fila_cab + 1},
                    },
                    "fields": "gridProperties.frozenRowCount",
                }
            }
        )
        px = anchos.get(titulo, [120] * cols)
        for i in range(cols):
            pedidos.append(
                {
                    "updateDimensionProperties": {
                        "range": {
                            "sheetId": sid,
                            "dimension": "COLUMNS",
                            "startIndex": i,
                            "endIndex": i + 1,
                        },
                        "properties": {"pixelSize": px[i] if i < len(px) else 110},
                        "fields": "pixelSize",
                    }
                }
            )
    if pedidos:
        sheets.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id, body={"requests": pedidos}
        ).execute()


def _compartir(drive, file_id: str, correos: list[str]) -> None:
    """# Aquí se da permiso de lectura a cada correo de CORREOS_AVISO."""
    for correo in correos:
        try:
            drive.permissions().create(
                fileId=file_id,
                body={
                    "type": "user",
                    "role": "reader",
                    "emailAddress": correo,
                },
                sendNotificationEmail=False,
                supportsAllDrives=True,
            ).execute()
        except HttpError as e:
            if e.resp.status in (400, 403):
                logger.info("Permiso de hoja ya existía o no se pudo añadir para %s", correo)
            else:
                raise


def publicar_xlsx_en_sheets(
    creds: Credentials,
    xlsx: Path,
    base: Path,
    correos: list[str],
) -> str:
    """
    # Aquí está el flujo principal:
      1) Abrir Excel
      2) Crear Sheet si no hay ID guardado (o reutilizar)
      3) Limpiar y pegar cada pestaña
      4) Formato + compartir
      5) Devolver el link (para el correo)
    """
    wb = load_workbook(xlsx, data_only=True)
    titulos = [ws.title for ws in wb.worksheets]
    sheets = build("sheets", "v4", credentials=creds, static_discovery=True)
    drive = build("drive", "v3", credentials=creds, static_discovery=True)

    sheet_id = _leer_id(base)
    if not sheet_id:
        creado = sheets.spreadsheets().create(
            body={
                "properties": {"title": TITULO_HOJA},
                "sheets": [{"properties": {"title": t}} for t in titulos],
            },
            fields="spreadsheetId,spreadsheetUrl",
        ).execute()
        sheet_id = creado["spreadsheetId"]
        _guardar_id(base, sheet_id)
        print(f"Hoja de inventario creada (se reutilizará): {creado.get('spreadsheetUrl')}", flush=True)
    else:
        _asegurar_pestanas(sheets, sheet_id, titulos)

    n_cols: dict[str, int] = {}
    for ws in wb.worksheets:
        valores = _filas_hoja(ws)
        n_cols[ws.title] = max((len(r) for r in valores), default=2)
        rango = f"'{ws.title}'"
        sheets.spreadsheets().values().clear(
            spreadsheetId=sheet_id, range=rango
        ).execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sheet_id,
            range=f"{rango}!A1",
            valueInputOption="RAW",
            body={"values": valores},
        ).execute()

    _aplicar_formato(sheets, sheet_id, titulos, n_cols)
    wb.close()
    _compartir(drive, sheet_id, correos)
    enlace = f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit"
    print(f"Inventario actualizado en Google Sheets: {enlace}", flush=True)
    return enlace
