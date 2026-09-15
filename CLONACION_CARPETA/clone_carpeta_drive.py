"""
CLONACION_CARPETA — clone_carpeta_drive.py
------------------------------------------
Entrada principal del bloque de clonación.

Qué hace, en orden:
  1. Lee RUTAS.xlsx (etiqueta, origen, destino).
  2. Clona cada fila origen → destino en Google Drive (reanudable).
  3. Genera inventario Excel local.
  4. Publica/sobrescribe una Google Sheet fija.
  5. Envía correo Gmail a CORREOS_AVISO con el link de la Sheet
     (si Sheets falla, adjunta el Excel).

No carga Cloud SQL (eso es LMS_Fabrica).

Credenciales: credentials.json (OAuth) + token.json (sesión).
"""
# ---------------------------------------------------------------------------
# IMPORTS del bloque de clonación
# ---------------------------------------------------------------------------

# Anotaciones de tipos modernas (str | None, list[...]).
from __future__ import annotations

# Log a archivo clonacion.log + consola.
import logging

# Errores HTTP de bajo nivel (reintentos de red).
import http.client

# Leer/escribir token.json (JSON).
import json

# Aleatorio para esperas (backoff) ante fallos temporales.
import random

# Extraer IDs de URLs de Drive con expresiones regulares.
import re

# Errores de red (DNS, socket, SSL) para reintentar.
import socket
import ssl
import sys
import time

# Contar nombres (duplicados) y mapas con valor por defecto.
from collections import Counter, defaultdict

# Estructura simple para una fila del Excel (etiqueta + ids).
from dataclasses import dataclass

# Rutas de archivos independientes del SO.
from pathlib import Path

# Si el refresh del token falla, hay que volver a loguear.
from google.auth.exceptions import RefreshError

# Renovar token vencido.
from google.auth.transport.requests import Request

# Credenciales OAuth de usuario.
from google.oauth2.credentials import Credentials

# Login de aplicación de escritorio (credentials.json).
from google_auth_oauthlib.flow import InstalledAppFlow

# Cliente API Drive.
from googleapiclient.discovery import build

# Errores HTTP de la API de Google.
from googleapiclient.errors import HttpError

# Argumentos de línea de comandos (--excel).
import argparse

# Leer el Excel RUTAS.xlsx.
from openpyxl import load_workbook

# Pasos 3–5: se importan aquí porque main() los encadena al final.
#   3) Inventario Excel local
#   4) Publicar a Google Sheets
#   5) Enviar correo Gmail
from reporte_inventario_clon import generar_reporte_excel
from notificar_clonacion import cargar_destinatarios, enviar_aviso_clonacion
from publicar_inventario_sheets import publicar_xlsx_en_sheets

# Ajustar si el Excel vive en otra ruta / otro PC.
RUTAS_XLSX = Path(
    r"C:\Users\angie_vera\Downloads\RUTAS.xlsx"
)
# Patrón para sacar el ID de una URL .../folders/ID
_DRIVE_FOLDER_RE = re.compile(r"/folders/([a-zA-Z0-9_-]+)")

# Permisos: Drive (clonar), Sheets (inventario), Gmail (aviso).
SCOPES = [
    "https://www.googleapis.com/auth/drive",
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/gmail.send",
]
MIME_FOLDER = "application/vnd.google-apps.folder"
BASE = Path(__file__).resolve().parent
TOKEN_PATH = BASE / "token.json"
CREDENTIALS_PATH = BASE / "credentials.json"


@dataclass(frozen=True)
class RutaClonacion:
    """Una fila del Excel ya normalizada a IDs de carpeta Drive."""

    etiqueta: str  # apodo humano del lote
    origen_id: str
    destino_id: str


def extraer_id_carpeta_drive(valor: str) -> str | None:
    """
    Acepta URL tipo .../folders/ID o el ID suelto.
    Devuelve el ID o None si el texto no es usable.
    """
    texto = valor.strip()
    if not texto:
        return None
    coincidencia = _DRIVE_FOLDER_RE.search(texto)
    if coincidencia:
        return coincidencia.group(1)
    if re.fullmatch(r"[a-zA-Z0-9_-]+", texto):
        return texto
    return None


def cargar_rutas_desde_excel(ruta: Path) -> list[RutaClonacion]:
    """
    Lee el Excel de rutas.
    Soporta 3 columnas (etiqueta|origen|destino) o 4 (cliente|etiqueta|origen|destino).
    En el clon, la columna cliente no decide la copia; solo se usan etiqueta/origen/destino.
    """
    if not ruta.is_file():
        print(f"No se encontró el archivo de rutas: {ruta}", file=sys.stderr)
        sys.exit(1)

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except PermissionError:
        print(
            f"No se pudo leer {ruta}. Cierra el Excel si lo tienes abierto e intenta de nuevo.",
            file=sys.stderr,
        )
        sys.exit(1)

    hoja = libro.active
    encabezados = [
        str(c).strip().lower() if c is not None else ""
        for c in next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
    ]

    def _indice(nombre: str, fallback: int) -> int:
        try:
            return encabezados.index(nombre)
        except ValueError:
            return fallback

    # Formato actual (3 cols): etiqueta | origen | destino
    # Formato nuevo (4 cols): cliente | etiqueta | origen | destino
    i_etiqueta = _indice("etiqueta", 0)
    i_origen = _indice("origen", 1)
    i_destino = _indice("destino", 2)
    if "origen" not in encabezados and "destino" not in encabezados:
        if len(encabezados) >= 4:
            i_etiqueta, i_origen, i_destino = 1, 2, 3

    rutas: list[RutaClonacion] = []
    for numero_fila, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        if not fila:
            continue
        n = len(fila)
        if max(i_etiqueta, i_origen, i_destino) >= n:
            continue

        etiqueta = str(fila[i_etiqueta]).strip() if fila[i_etiqueta] is not None else ""
        origen_raw = str(fila[i_origen]).strip() if fila[i_origen] is not None else ""
        destino_raw = str(fila[i_destino]).strip() if fila[i_destino] is not None else ""

        if not origen_raw and not destino_raw:
            continue

        origen_id = extraer_id_carpeta_drive(origen_raw)
        destino_id = extraer_id_carpeta_drive(destino_raw)

        if not origen_id or not destino_id:
            print(
                f"Fila {numero_fila}: URLs o IDs inválidos "
                f"(origen={origen_raw!r}, destino={destino_raw!r}).",
                file=sys.stderr,
            )
            sys.exit(1)

        if not etiqueta:
            etiqueta = f"Fila {numero_fila}"

        rutas.append(RutaClonacion(etiqueta=etiqueta, origen_id=origen_id, destino_id=destino_id))

    libro.close()

    if not rutas:
        print(f"No hay rutas válidas en {ruta}.", file=sys.stderr)
        sys.exit(1)

    return rutas


def _oauth_escritorio() -> Credentials:
    """Abre flujo OAuth en el navegador (URL manual) y guarda token.json."""
    if not CREDENTIALS_PATH.is_file():
        print(
            f"Token inválido o ausente y falta {CREDENTIALS_PATH} para volver a autorizar.",
            file=sys.stderr,
        )
        sys.exit(1)
    print(
        "No se abre el navegador solo. Copie la URL que aparece abajo y péguela "
        "en el navegador donde YA ve DESTINO-PRUEBA (cuenta fábrica de contenidos).",
        flush=True,
    )
    flow = InstalledAppFlow.from_client_secrets_file(str(CREDENTIALS_PATH), SCOPES)
    creds = flow.run_local_server(
        port=0,
        open_browser=False,
        authorization_prompt_message="Abra esta URL en ese navegador:\n{url}\n",
        success_message="Listo. Ya puede volver a la terminal. Esta pestaña se puede cerrar.",
    )
    TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")
    print(f"Credenciales guardadas en {TOKEN_PATH}", flush=True)
    return creds


def cargar_credenciales() -> Credentials:
    """
    Orden de autenticación:
      1) Usar token.json si es válido y tiene todos los SCOPES (incluye Gmail).
      2) Si venció, intentar refresh.
      3) Si no, login de escritorio con credentials.json.
    """
    creds: Credentials | None = None
    if TOKEN_PATH.is_file():
        try:
            with TOKEN_PATH.open(encoding="utf-8") as f:
                data = json.load(f)
            creds = Credentials.from_authorized_user_info(data, SCOPES)
        except (json.JSONDecodeError, OSError, ValueError) as e:
            print(f"Aviso: no se pudo leer token.json ({e}); se pedirá login.", flush=True)
            creds = None

    if creds and creds.valid:
        if creds.scopes and not set(SCOPES) <= set(creds.scopes):
            print(
                "El token no incluye permiso de correo. Hay que autorizar de nuevo en el navegador…",
                flush=True,
            )
            return _oauth_escritorio()
        return creds

    if creds and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
            TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")
            return creds
        except RefreshError as e:
            print(
                f"El refresh del token falló ({e}); se iniciará login de nuevo.",
                flush=True,
            )

    return _oauth_escritorio()


def ejecutar(req, max_intentos: int = 12, num_retries_http: int = 5):
    """Ejecuta una petición de la API con reintentos ante fallos de red transitorios."""
    ultimo: Exception | None = None
    for intento in range(max_intentos):
        try:
            return req.execute(num_retries=num_retries_http)
        except HttpError as e:
            st = e.resp.status
            if st in (400, 401, 403, 404):
                raise
            if st not in (408, 429) and not (500 <= st <= 599):
                raise
            ultimo = e
        except (
            http.client.IncompleteRead,
            http.client.RemoteDisconnected,
            socket.gaierror,
            socket.error,
            ConnectionError,
            OSError,
            ssl.SSLError,
        ) as e:
            ultimo = e
        if intento < max_intentos - 1:
            espera = min(2**intento * (0.4 + random.random() * 0.4), 90.0)
            time.sleep(espera)
    if ultimo is not None:
        raise ultimo
    raise RuntimeError("Error desconocido al ejecutar la petición")


def listar_hijos(svc, parent_id: str) -> list[dict]:
    """
    Lista TODOS los hijos directos de una carpeta Drive (no recursivo).

    - Páginas de hasta 1000 hasta agotar nextPageToken.
    - Solo id, name, mimeType (bastan para clonar / inventariar).
    - Incluye Shared Drives; ignora papelera.
    """
    out: list[dict] = []
    page = None
    while True:
        # Query: hijos de este padre y no en papelera.
        q = f"'{parent_id}' in parents and trashed = false"
        r = ejecutar(
            svc.files()
            .list(
                q=q,
                spaces="drive",
                fields="nextPageToken, files(id, name, mimeType)",
                pageToken=page,
                pageSize=1000,
                includeItemsFromAllDrives=True,
                supportsAllDrives=True,
            )
        )
        out.extend(r.get("files", []))
        page = r.get("nextPageToken")
        if not page:
            break
    return out


def indice_destino(
    svc, parent_id: str
) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """
    Índice en memoria de lo que YA hay en una carpeta destino.

    Devuelve dos mapas nombre → lista de IDs:
      - carpetas: para reanudar (entrar a la existente) o detectar duplicados
      - archivos: para omitir si ya está, o contar cuántos hay del mismo nombre

    Varios IDs con el mismo nombre = posibles duplicados de corridas previas.
    """
    carpetas: dict[str, list[str]] = {}
    archivos: dict[str, list[str]] = {}
    for h in listar_hijos(svc, parent_id):
        n, mid = h["name"], h["mimeType"]
        if mid == MIME_FOLDER:
            carpetas.setdefault(n, []).append(h["id"])
        else:
            archivos.setdefault(n, []).append(h["id"])
    return carpetas, archivos


def _esperar_indice(
    svc,
    parent_id: str,
    min_archivos: int,
    min_carpetas: int,
    intentos: int = 8,
) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """
    Drive a veces tarda en indexar lo recién copiado.
    Relee el destino hasta ver al menos min_archivos / min_carpetas
    (o agotar intentos). Evita crear/copiar de más por un listado atrasado.
    """
    ultimo: tuple[dict[str, list[str]], dict[str, list[str]]] | None = None
    for i in range(intentos):
        carpetas, archivos = indice_destino(svc, parent_id)
        n_arc = sum(len(v) for v in archivos.values())
        n_car = sum(len(v) for v in carpetas.values())
        ultimo = (carpetas, archivos)
        if n_arc >= min_archivos and n_car >= min_carpetas:
            return ultimo
        if i < intentos - 1:
            time.sleep(1.5)
    return ultimo if ultimo is not None else indice_destino(svc, parent_id)


def _enviar_a_papelera(svc, file_id: str) -> None:
    """
    Marca un archivo/carpeta como trashed (papelera), sin borrado permanente.
    404/410 se ignoran (ya no está); otros errores HTTP se relanzan.
    """
    try:
        ejecutar(
            svc.files().update(
                fileId=file_id,
                body={"trashed": True},
                supportsAllDrives=True,
            ),
            num_retries_http=0,
        )
    except HttpError as e:
        if e.resp.status not in (404, 410):
            raise


def _quitar_duplicados_en_carpeta(
    svc,
    destino_parent_id: str,
    cupo_archivos: Counter,
    cupo_carpetas: Counter,
    dest_carpetas: dict[str, list[str]],
    dest_archivos: dict[str, list[str]],
    contador: dict[str, int],
) -> None:
    """
    Deja en destino exactamente las cantidades del origen (cupo_*).

    Flujo:
      1) Refresca el índice real de Drive en dest_carpetas / dest_archivos.
      2) Por cada nombre: conserva los primeros ``keep`` IDs; el resto → papelera.
      3) Actualiza el contador de duplicados quitados.

    Así una reanudación no acumula copias de más del mismo nombre.
    """
    # Refresco: lo que Drive muestra ahora (puede haber cambiado tras copias).
    ahora_c, ahora_a = indice_destino(svc, destino_parent_id)
    dest_carpetas.clear()
    dest_carpetas.update(ahora_c)
    dest_archivos.clear()
    dest_archivos.update(ahora_a)
    for nombre, ids in list(dest_carpetas.items()):
        keep = cupo_carpetas.get(nombre, 0)
        extras = ids[keep:]
        if extras:
            dest_carpetas[nombre] = ids[:keep]
        for fid in extras:
            print(f"  [quitar carpeta duplicada] {nombre}", flush=True)
            _enviar_a_papelera(svc, fid)
            contador["carpetas_duplicadas_quitadas"] = (
                contador.get("carpetas_duplicadas_quitadas", 0) + 1
            )
    for nombre, ids in list(dest_archivos.items()):
        keep = cupo_archivos.get(nombre, 0)
        extras = ids[keep:]
        if extras:
            dest_archivos[nombre] = ids[:keep]
        for fid in extras:
            print(f"  [quitar duplicado] {nombre}", flush=True)
            _enviar_a_papelera(svc, fid)
            contador["archivos_duplicados_quitados"] = (
                contador.get("archivos_duplicados_quitados", 0) + 1
            )


def _copiar_archivo_sin_duplicar(
    svc,
    origen_file_id: str,
    nombre: str,
    destino_parent_id: str,
    dest_archivos: dict[str, list[str]],
    contador: dict[str, int],
) -> None:
    """
    Copia un archivo al destino evitando duplicados tras timeouts.

    Idea clave: si la API falla por red/timeout, Drive a veces YA creó
    la copia. Antes de reintentar se relee el índice; si la cantidad con
    ese nombre aumentó, se cuenta éxito y NO se vuelve a copiar.
    Errores duros (400/401/403/404) → fallido sin reintentar.
    """
    # Cuántos había con este nombre antes de intentar (para detectar “ya está”).
    antes = len(dest_archivos.get(nombre) or [])
    print(f"  [archivo] {nombre}", flush=True)
    ultimo: Exception | None = None
    for intento in range(12):
        try:
            creado = (
                svc.files()
                .copy(
                    fileId=origen_file_id,
                    body={"name": nombre, "parents": [destino_parent_id]},
                    fields="id",
                    supportsAllDrives=True,
                )
                .execute(num_retries=0)
            )
            nid = creado.get("id")
            dest_archivos.setdefault(nombre, []).append(nid or "ok")
            contador["archivos_copiados"] = contador.get("archivos_copiados", 0) + 1
            return
        except HttpError as e:
            st = e.resp.status
            if st in (400, 401, 403, 404):
                print(f"  [ERROR] no se pudo copiar {nombre}: {e}", flush=True)
                contador["archivos_fallidos"] = contador.get("archivos_fallidos", 0) + 1
                return
            if st not in (408, 429) and not (500 <= st <= 599):
                print(f"  [ERROR] no se pudo copiar {nombre}: {e}", flush=True)
                contador["archivos_fallidos"] = contador.get("archivos_fallidos", 0) + 1
                return
            ultimo = e
        except (
            http.client.IncompleteRead,
            http.client.RemoteDisconnected,
            socket.timeout,
            socket.gaierror,
            socket.error,
            ConnectionError,
            OSError,
            ssl.SSLError,
            TimeoutError,
        ) as e:
            ultimo = e
        # Tras fallo temporal: ¿Drive ya tiene el archivo? Entonces no duplicar.
        _, ahora = indice_destino(svc, destino_parent_id)
        dest_archivos.clear()
        dest_archivos.update(ahora)
        if len(dest_archivos.get(nombre) or []) > antes:
            print(
                f"  [archivo] {nombre} (Drive ya lo tenía tras el timeout; no duplico)",
                flush=True,
            )
            contador["archivos_copiados"] = contador.get("archivos_copiados", 0) + 1
            return
        if intento < 11:
            time.sleep(min(2**intento * (0.4 + random.random() * 0.4), 90.0))
    print(f"  [ERROR] no se pudo copiar {nombre}: {ultimo}", flush=True)
    contador["archivos_fallidos"] = contador.get("archivos_fallidos", 0) + 1


def copiar_arbol(
    svc,
    origen_id: str,
    destino_parent_id: str,
    contador: dict[str, int],
) -> None:
    """
    Copia recursiva origen → destino.
    - Si la subcarpeta ya existe en destino: reanuda dentro de ella.
    - Si el archivo ya existe (mismo nombre): omite.
    - Si falta: crea carpeta o copia archivo.
    Al final limpia duplicados de más respecto al origen.
    """
    origen_hijos = listar_hijos(svc, origen_id)
    dest_carpetas, dest_archivos = indice_destino(svc, destino_parent_id)
    cupo_a = Counter(
        h["name"] for h in origen_hijos if h.get("mimeType") != MIME_FOLDER
    )
    cupo_c = Counter(
        h["name"] for h in origen_hijos if h.get("mimeType") == MIME_FOLDER
    )
    visto: dict[str, int] = defaultdict(int)
    for item in origen_hijos:
        nombre = item["name"]
        mid = item["mimeType"]
        if mid == MIME_FOLDER:
            existentes = dest_carpetas.get(nombre) or []
            if existentes:
                print(f"  [carpeta] {nombre} (reanudar)", flush=True)
                contador["carpetas_reusadas"] = contador.get("carpetas_reusadas", 0) + 1
                copiar_arbol(svc, item["id"], existentes[0], contador)
            else:
                print(f"  [carpeta] {nombre}", flush=True)
                creada = ejecutar(
                    svc.files()
                    .create(
                        body={
                            "name": nombre,
                            "parents": [destino_parent_id],
                            "mimeType": MIME_FOLDER,
                        },
                        fields="id",
                        supportsAllDrives=True,
                    )
                )
                dest_carpetas.setdefault(nombre, []).append(creada["id"])
                contador["carpetas_nuevas"] = contador.get("carpetas_nuevas", 0) + 1
                copiar_arbol(svc, item["id"], creada["id"], contador)
        else:
            if visto[nombre] < len(dest_archivos.get(nombre) or []):
                print(f"  [archivo] {nombre} (ya existe, omito)", flush=True)
                contador["archivos_omitidos"] = contador.get("archivos_omitidos", 0) + 1
                visto[nombre] += 1
                continue
            _copiar_archivo_sin_duplicar(
                svc, item["id"], nombre, destino_parent_id, dest_archivos, contador
            )
            visto[nombre] += 1
    _quitar_duplicados_en_carpeta(
        svc,
        destino_parent_id,
        cupo_a,
        cupo_c,
        dest_carpetas,
        dest_archivos,
        contador,
    )


def igualar_arbol(
    svc,
    origen_id: str,
    destino_parent_id: str,
    contador: dict[str, int],
) -> None:
    """
    Segunda pasada tras copiar_arbol.
    Espera a que Drive indexe el destino, quita duplicados de más y completa
    archivos/carpetas que aún falten respecto al origen.
    """
    origen_hijos = listar_hijos(svc, origen_id)
    n_ori_a = sum(1 for h in origen_hijos if h.get("mimeType") != MIME_FOLDER)
    n_ori_c = sum(1 for h in origen_hijos if h.get("mimeType") == MIME_FOLDER)
    dest_carpetas, dest_archivos = _esperar_indice(
        svc, destino_parent_id, n_ori_a, n_ori_c
    )
    cupo_a = Counter(
        h["name"] for h in origen_hijos if h.get("mimeType") != MIME_FOLDER
    )
    cupo_c = Counter(
        h["name"] for h in origen_hijos if h.get("mimeType") == MIME_FOLDER
    )
    _quitar_duplicados_en_carpeta(
        svc,
        destino_parent_id,
        cupo_a,
        cupo_c,
        dest_carpetas,
        dest_archivos,
        contador,
    )
    usados: dict[str, int] = defaultdict(int)
    for item in origen_hijos:
        nombre = item["name"]
        if item.get("mimeType") == MIME_FOLDER:
            continue
        if usados[nombre] < len(dest_archivos.get(nombre) or []):
            usados[nombre] += 1
            continue
        print(f"  [completar archivo] {nombre}", flush=True)
        _copiar_archivo_sin_duplicar(
            svc, item["id"], nombre, destino_parent_id, dest_archivos, contador
        )
        usados[nombre] += 1
    _quitar_duplicados_en_carpeta(
        svc,
        destino_parent_id,
        cupo_a,
        cupo_c,
        dest_carpetas,
        dest_archivos,
        contador,
    )
    for item in origen_hijos:
        if item.get("mimeType") != MIME_FOLDER:
            continue
        nombre = item["name"]
        existentes = dest_carpetas.get(nombre) or []
        if existentes:
            ex_id = existentes[0]
        else:
            print(f"  [completar carpeta] {nombre}", flush=True)
            creada = ejecutar(
                svc.files()
                .create(
                    body={
                        "name": nombre,
                        "parents": [destino_parent_id],
                        "mimeType": MIME_FOLDER,
                    },
                    fields="id",
                    supportsAllDrives=True,
                )
            )
            ex_id = creada["id"]
            dest_carpetas.setdefault(nombre, []).append(ex_id)
            contador["carpetas_nuevas"] = contador.get("carpetas_nuevas", 0) + 1
        igualar_arbol(svc, item["id"], ex_id, contador)


def _diagnostico_404(svc, folder_id: str) -> None:
    """
    Ayuda cuando files().get da 404: muestra el Gmail del token y busca
    carpetas llamadas DESTINO-PRUEBA visibles para esa cuenta.
    Suele indicar ID mal pegado en el Excel o login con otra cuenta.
    """
    yo = ejecutar(svc.about().get(fields="user(emailAddress)"))
    correo = (yo.get("user") or {}).get("emailAddress", "?")
    print(f"  Gmail del token: {correo}", file=sys.stderr, flush=True)
    print(f"  ID que no encontró: {folder_id}", file=sys.stderr, flush=True)
    r = ejecutar(
        svc.files()
        .list(
            q=(
                "name = 'DESTINO-PRUEBA' and mimeType = "
                f"'{MIME_FOLDER}' and trashed = false"
            ),
            corpora="allDrives",
            includeItemsFromAllDrives=True,
            supportsAllDrives=True,
            fields="files(id, name)",
            pageSize=10,
        )
    )
    halladas = r.get("files") or []
    if halladas:
        print("  Esta cuenta SÍ tiene carpeta(s) DESTINO-PRUEBA:", file=sys.stderr, flush=True)
        for f in halladas:
            print(f"    - {f['name']}  id={f['id']}", file=sys.stderr, flush=True)
        print(
            "  Si el id no coincide con el del Excel, use el id de arriba en RUTAS.xlsx.",
            file=sys.stderr,
            flush=True,
        )
    else:
        print(
            "  Esta cuenta NO tiene ninguna carpeta llamada DESTINO-PRUEBA. "
            "En el login de Google elija el usuario que en Drive es dueño de esa carpeta.",
            file=sys.stderr,
            flush=True,
        )


def _validar_carpeta(svc, folder_id: str, etiqueta: str) -> dict:
    """
    Comprueba que el ID existe y es una carpeta Drive.
    Si falla con 404, llama a _diagnostico_404 y sale del proceso.
    Devuelve el dict con id, name, mimeType.
    """
    try:
        carpeta = ejecutar(
            svc.files().get(
                fileId=folder_id,
                fields="id, name, mimeType",
                supportsAllDrives=True,
            )
        )
    except HttpError as e:
        print(f"Error al leer {etiqueta}: {e}", file=sys.stderr)
        if getattr(e.resp, "status", None) == 404:
            _diagnostico_404(svc, folder_id)
        sys.exit(1)

    if carpeta.get("mimeType") != MIME_FOLDER:
        print(f"El ID de {etiqueta} no es una carpeta.", file=sys.stderr)
        sys.exit(1)
    return carpeta


def main(argv: list[str] | None = None) -> None:
    """
    Orquestación completa del bloque de clonación:
      Excel → clonar cada ruta → inventario → Sheets → correo.
    """
    parser = argparse.ArgumentParser(
        description="Clona carpetas Drive según RUTAS.xlsx y encadena inventario + correo."
    )
    parser.add_argument(
        "--excel",
        type=Path,
        default=RUTAS_XLSX,
        help=f"Ruta a RUTAS.xlsx (default: {RUTAS_XLSX})",
    )
    args = parser.parse_args(argv)
    excel_rutas = args.excel.resolve()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(BASE / "clonacion.log", encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )
    rutas = cargar_rutas_desde_excel(excel_rutas)
    print(f"Iniciando (API Google Drive) — {len(rutas)} rutas en {excel_rutas.name}…", flush=True)
    creds = cargar_credenciales()
    svc = build("drive", "v3", credentials=creds, static_discovery=True)
    yo = ejecutar(svc.about().get(fields="user(emailAddress,displayName)"))
    correo_token = (yo.get("user") or {}).get("emailAddress", "(desconocido)")
    print(f"Cuenta de Google que usa el script: {correo_token}", flush=True)
    print(
        "Si no es la misma con la que ves DESTINO-PRUEBA en el navegador, "
        "borra token.json y vuelve a autorizar eligiendo esa cuenta.",
        flush=True,
    )

    estado_correo = "exitoso"
    nota_error = ""
    generado: Path | None = None
    resumenes: list[dict] = []

    try:
        contador: dict[str, int] = {}
        total = len(rutas)
        trabajos_reporte: list[dict] = []

        for indice, ruta in enumerate(rutas, start=1):
            origen = _validar_carpeta(svc, ruta.origen_id, f"origen [{ruta.etiqueta}]")
            destino = _validar_carpeta(svc, ruta.destino_id, f"destino [{ruta.etiqueta}]")
            print(
                f"[{indice}/{total}] {ruta.etiqueta}: "
                f"«{origen['name']}» → «{destino['name']}»…",
                flush=True,
            )
            copiar_arbol(svc, ruta.origen_id, ruta.destino_id, contador)
            print("Igualando con el origen (huecos y duplicados)…", flush=True)
            igualar_arbol(svc, ruta.origen_id, ruta.destino_id, contador)
            trabajos_reporte.append(
                {
                    "etiqueta": ruta.etiqueta,
                    "origen_id": ruta.origen_id,
                    "destino_id": ruta.destino_id,
                    "origen_nombre": origen["name"],
                    "destino_nombre": destino["name"],
                }
            )
            if indice < total:
                print(f"[{indice}/{total}] {ruta.etiqueta} completado.", flush=True)

        print("Listo.", flush=True)
        print(f"  Rutas procesadas: {total}", flush=True)
        print(f"  Carpetas nuevas: {contador.get('carpetas_nuevas', 0)}", flush=True)
        print(
            f"  Carpetas reanudadas (ya existían): {contador.get('carpetas_reusadas', 0)}",
            flush=True,
        )
        print(f"  Archivos copiados: {contador.get('archivos_copiados', 0)}", flush=True)
        print(
            f"  Archivos ya existentes (omitidos): {contador.get('archivos_omitidos', 0)}",
            flush=True,
        )
        if contador.get("archivos_fallidos"):
            print(
                f"  Archivos que no se pudieron copiar: {contador['archivos_fallidos']}",
                flush=True,
            )
        if contador.get("carpetas_duplicadas_quitadas"):
            print(
                f"  Carpetas duplicadas a la papelera: {contador['carpetas_duplicadas_quitadas']}",
                flush=True,
            )
        if contador.get("archivos_duplicados_quitados"):
            print(
                f"  Duplicados enviados a la papelera: {contador['archivos_duplicados_quitados']}",
                flush=True,
            )

        ruta_xlsx = BASE / "reporte_inventario_clon.xlsx"
        print(
            "Actualizando inventario (se sobrescribe el mismo archivo y la Google Sheet)…",
            flush=True,
        )
        # Paso 3: inventario local (Excel con hojas Resumen, Faltantes, etc.).
        generado, resumenes = generar_reporte_excel(
            svc, listar_hijos, trabajos_reporte, ruta_xlsx
        )
        print(f"  Inventario local: {generado}", flush=True)
        if resumenes and not all(r.get("validacion_ok") for r in resumenes):
            estado_correo = "con diferencias"
    except Exception as e:
        estado_correo = "con error"
        nota_error = str(e)
        logging.exception("Fallo durante la clonación o el inventario")
        print(f"Error: {e}", file=sys.stderr)

    destinatarios = cargar_destinatarios(BASE)
    enlace_hoja = ""
    if generado is not None and generado.is_file():
        try:
            # Paso 4: misma info del Excel → una Sheet fija (se sobrescribe).
            enlace_hoja = publicar_xlsx_en_sheets(creds, generado, BASE, destinatarios)
        except Exception:
            logging.exception("No se pudo actualizar Google Sheets; se adjuntará el Excel")
            print(
                "No se pudo publicar en Google Sheets (active la API de Sheets en Google Cloud). "
                "El correo llevará el Excel adjunto.",
                flush=True,
            )
    try:
        # Paso 5: correo Gmail con link de Sheet (o Excel adjunto si falló Sheets).
        adjunto = None if enlace_hoja else generado
        enviar_aviso_clonacion(
            creds,
            destinatarios,
            estado_correo,
            resumenes,
            adjunto,
            nota_error,
            enlace_hoja,
        )
    except Exception as e:
        logging.exception("No se pudo enviar el correo de aviso")
        print(
            "No se pudo enviar el correo. En Google Cloud active la API de Gmail "
            f"y vuelva a autorizar. Detalle: {e}",
            file=sys.stderr,
        )

    if estado_correo == "con error":
        sys.exit(1)


if __name__ == "__main__":
    main()
